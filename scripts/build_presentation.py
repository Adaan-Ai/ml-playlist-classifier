#!/usr/bin/env python3
"""Build the 8-slide ML mini-project deck (16:9) with python-pptx.

Run from anywhere:  python scripts/build_presentation.py
Output:             output/project_presentation_cv_final.pptx

Metrics come from results/metrics.json when it has this shape, otherwise the
hardcoded benchmark values below are used:

{
  "cv":   {"Logistic Regression": {"micro_f1": .., "macro_f1": .., "top1": ..,
                                   "recall3": .., "subset": ..},
           "RBF SVM": {...}, "Random Forest": {...}},
  "test": {"macro_f1": .., "micro_f1": .., "subset": .., "top1": .., "recall3": ..},
  "per_genre": [["Rock", 0.703, 896], ...]
}
All ratios are fractions (0.5487, not 54.87).
"""
import io
import json
import re
import textwrap
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
METRICS_PATH = ROOT / "results" / "metrics.json"
OUT_PATH = ROOT / "output" / "project_presentation_cv_final.pptx"

# ---------------------------------------------------------------- palette
NAVY = "15233B"
TEAL = "176B67"
MINT = "DDF3ED"
CHARCOAL = "1F2937"
SLATE = "526078"
CARD = "F5F7FA"
BORDER = "D8DEE8"
WHITE = "FFFFFF"

FONT = "Calibri"
MONO = "Courier New"
SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
MARGIN = 0.5
CONTENT_W = SLIDE_W_IN - 2 * MARGIN  # 12.333
TOTAL_SLIDES = 8
COURSE_TAG = "UE24CS352A  ·  ML Mini-Project"

WARNINGS = []

# ---------------------------------------------------------------- data
DEFAULT_METRICS = {
    "cv": {
        "Logistic Regression": {"micro_f1": 0.463, "macro_f1": 0.473, "top1": 0.485, "recall3": 0.826, "subset": 0.097},
        "RBF SVM": {"micro_f1": 0.540, "macro_f1": 0.528, "top1": 0.562, "recall3": 0.864, "subset": 0.482},
        "Random Forest": {"micro_f1": 0.570, "macro_f1": 0.556, "top1": 0.590, "recall3": 0.879, "subset": 0.541},
    },
    "test": {"macro_f1": 0.563, "micro_f1": 0.576, "subset": 0.5487, "top1": 0.5936, "recall3": 0.8815},
    "per_genre": [
        ["Rock", 0.703, 896],
        ["EDM", 0.683, 1158],
        ["Rap", 0.653, 1151],
        ["R&B", 0.491, 1023],
        ["Latin", 0.457, 879],
        ["Pop", 0.391, 948],
    ],
}
CV_KEYS = ("micro_f1", "macro_f1", "top1", "recall3", "subset")
MODELS = ("Logistic Regression", "RBF SVM", "Random Forest")


def load_metrics():
    """Return metrics from results/metrics.json if valid, else the defaults."""
    if not METRICS_PATH.exists():
        print(f"[info] {METRICS_PATH.name} not found, using hardcoded benchmark values.")
        return DEFAULT_METRICS
    try:
        raw = json.loads(METRICS_PATH.read_text())
        # Native presentation schema
        if "cv" in raw and "test" in raw and "per_genre" in raw:
            return raw
        # Map from training output schema (results/metrics.json)
        models_data = raw.get("models", {})
        test_eval = raw.get("test_evaluation", {})
        cv = {}
        for m in MODELS:
            mod = models_data.get(m, {})
            cv[m] = {
                "micro_f1": mod.get("micro_f1", 0.0),
                "macro_f1": mod.get("macro_f1", 0.0),
                "top1": mod.get("precision_at_1", 0.0),
                "recall3": mod.get("recall_at_3", 0.0),
                "subset": mod.get("subset_accuracy", 0.0),
            }
        test = {
            "macro_f1": test_eval.get("macro_f1", 0.0),
            "micro_f1": test_eval.get("micro_f1", 0.0),
            "subset": test_eval.get("subset_accuracy", 0.0),
            "top1": test_eval.get("precision_at_1", 0.0),
            "recall3": test_eval.get("recall_at_3", 0.0),
        }
        per_genre_raw = test_eval.get("per_genre", {})
        display = {"rock": "Rock", "edm": "EDM", "rap": "Rap", "r&b": "R&B", "latin": "Latin", "pop": "Pop"}
        per_genre = [
            [display.get(k, k.capitalize()), v.get("f1", 0.0), v.get("support", 0)]
            for k, v in sorted(per_genre_raw.items(), key=lambda x: x[1].get("f1", 0), reverse=True)
        ]
        print(f"[info] Loaded and mapped metrics from {METRICS_PATH}")
        return {"cv": cv, "test": test, "per_genre": per_genre}
    except Exception as exc:
        print(f"[warn] {METRICS_PATH.name} mapping error ({exc!r}); using hardcoded values.")
        return DEFAULT_METRICS


METRICS = load_metrics()


def f3(x):
    return f"{x:.3f}"


def pct(x, d=1):
    return f"{round(x * 100 + 1e-9, d):.{d}f}%"


# ---------------------------------------------------------------- low-level helpers
def _rgb(hex_str):
    return RGBColor.from_string(hex_str)


def _est_height_in(paras, width_in, default_size, space_after_pt=0, bullet_indent=0.0):
    """Rough text height estimate (inches) used to flag likely overflow."""
    total = 0.0
    for item in paras:
        size = item.get("size", default_size)
        bold = item.get("bold", False)
        text = item.get("lead", "") + item.get("text", "")
        indent = bullet_indent if item.get("bullet") else 0.0
        cpl = max(1, int((width_in - indent) * 72 / (size * (0.53 if bold else 0.5))))
        lines = max(1, len(textwrap.wrap(text, cpl)) if text else 1)
        total += lines * size * 1.2 / 72 + item.get("space_after", space_after_pt) / 72
    return total


def _norm_paras(paras):
    if isinstance(paras, str):
        paras = [paras]
    return [{"text": p} if isinstance(p, str) else dict(p) for p in paras]


def _set_bullet(p, indent_in=0.2):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in))))
    pPr.set("indent", str(-int(Inches(indent_in))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•")


def _text(slide, x, y, w, h, paras, size=14, color=CHARCOAL, bold=False,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT,
          space_after=0, label=""):
    """Textbox with fixed size, word wrap and no autofit. Warns if text likely overflows."""
    items = _norm_paras(paras)
    for it in items:
        it.setdefault("bold", bold)
    est = _est_height_in(items, w, size, space_after, bullet_indent=0.2)
    if est > h + 0.02:
        WARNINGS.append(f"possible overflow in '{label or items[0].get('text', '')[:30]}': need ~{est:.2f}in, have {h:.2f}in")

    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor

    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = it.get("align", align)
        p.space_after = Pt(it.get("space_after", space_after))
        p.line_spacing = 1.0
        runs = []
        if it.get("lead"):
            runs.append((it["lead"], True, it.get("lead_color", it.get("color", color))))
        runs.append((it.get("text", ""), it.get("bold", bold), it.get("color", color)))
        for txt, b, col in runs:
            r = p.add_run()
            r.text = txt
            r.font.name = it.get("font", font)
            r.font.size = Pt(it.get("size", size))
            r.font.bold = b
            r.font.color.rgb = _rgb(col)
        if it.get("bullet"):
            _set_bullet(p)
    return tb


def _rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, line_w=0.75):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    ref = shp._element.find('.//' + qn('a:effectRef'))
    if ref is not None:
        ref.set('idx', '0')
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    return shp


def _card(slide, x, y, w, h, fill=CARD, line=BORDER):
    """Rounded container used for every content block."""
    return _rect(slide, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)


def _badge(slide, x, y, w, h, text, fill=MINT, color=TEAL, size=11, bold=True,
           shape=MSO_SHAPE.ROUNDED_RECTANGLE, font=FONT):
    shp = _rect(slide, x, y, w, h, fill=fill, shape=shape, radius=0.5 if shape == MSO_SHAPE.ROUNDED_RECTANGLE else None)
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = _rgb(color)
    return shp


def _cell_border(cell, color=BORDER, w=9525):
    """Borders must precede the fill element in a:tcPr, so call this before setting fill."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag), w=str(w), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")


def _table(slide, x, y, col_widths, row_h, rows, size=14, header_fill=NAVY,
           highlight_rows=(), first_col_align=PP_ALIGN.LEFT, header_size=None, mono_cols=(), left_cols=()):
    """rows[0] is the header. row_h is a float or a list of heights (inches)."""
    n_rows, n_cols = len(rows), len(col_widths)
    heights = row_h if isinstance(row_h, (list, tuple)) else [row_h] * n_rows
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y),
                                   Inches(sum(col_widths)), Inches(sum(heights)))
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, w in enumerate(col_widths):
        tbl.columns[i].width = Inches(w)
    for i, h in enumerate(heights):
        tbl.rows[i].height = Inches(h)

    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = tbl.cell(r_i, c_i)
            is_header = r_i == 0
            highlighted = r_i in highlight_rows
            _cell_border(cell)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(header_fill if is_header else (MINT if highlighted else WHITE))
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = first_col_align if (c_i == 0 or c_i in left_cols) else PP_ALIGN.CENTER
            run = p.add_run()
            run.text = str(val)
            run.font.name = MONO if c_i in mono_cols and not is_header else FONT
            run.font.size = Pt((header_size or size) if is_header else size)
            run.font.bold = is_header or highlighted or c_i == 0
            run.font.color.rgb = _rgb(WHITE if is_header else (TEAL if highlighted else CHARCOAL))
    return shape


def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = textwrap.dedent(text).strip()


def _set_background(slide, hex_color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(hex_color)


def _chrome(slide, n, title, category):
    """Title, accent rule, category badge (top right) and footer (bottom right)."""
    _set_background(slide, WHITE)
    _text(slide, MARGIN, 0.38, 10.0, 0.8, title, size=26, color=NAVY, bold=True,
          anchor=MSO_ANCHOR.MIDDLE, label=f"title {n}")
    _rect(slide, MARGIN, 1.2, 0.9, 0.05, fill=TEAL)
    _badge(slide, 10.73, 0.55, 2.1, 0.4, category, size=11)
    _text(slide, 7.83, 7.02, 5.0, 0.28, f"{COURSE_TAG}  ·  Slide {n} / {TOTAL_SLIDES}",
          size=10, color=SLATE, align=PP_ALIGN.RIGHT, label=f"footer {n}")


def _number_badge(slide, x, y, n, d=0.46):
    return _badge(slide, x, y, d, d, str(n), fill=TEAL, color=WHITE, size=15, shape=MSO_SHAPE.OVAL)


# ---------------------------------------------------------------- slides
def slide_1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(s, NAVY)
    _rect(s, 0, 0, 0.35, SLIDE_H_IN, fill=TEAL)

    _badge(s, 1.0, 0.9, 4.9, 0.46, "UE24CS352A  ·  Machine Learning Mini-Project", fill=TEAL, color=WHITE, size=13)
    _text(s, 1.0, 1.75, 11.3, 1.8, "Training a Playlist Curator Based on Track Features",
          size=44, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, label="title slide")
    _text(s, 1.0, 3.8, 10.5, 0.9,
          "Multi-label genre prediction from 12 acoustic attributes, evaluated with track-level leakage control",
          size=20, color=MINT, label="subtitle")
    _rect(s, 1.0, 5.0, 2.2, 0.05, fill=TEAL)
    _text(s, 1.0, 5.3, 11.0, 0.5, "Mohammed Adaan Hamad  ·  Lakshya Jeet Singh",
          size=22, color=WHITE, bold=True, label="team")
    _text(s, 1.0, 5.9, 11.0, 0.4, "Section E  ·  PES University", size=15, color=BORDER, label="section")

    _notes(s, f"""
        Opening (about 45 seconds). We are Mohammed Adaan Hamad and Lakshya Jeet Singh from Section E, presenting our
        UE24CS352A Machine Learning mini-project: training a playlist curator based on track features.

        The question: given only the audio characteristics of a song, can a model say which playlist genres it belongs to?
        Our target is six genres (edm, latin, pop, r&b, rap, rock), and a song can carry more than one, so this is
        multi-label classification.

        Roadmap: problem and scope; the dataset; how we prevented leakage; the 5-fold cross-validation comparison of
        three models; the held-out test result for the winner (Random Forest); the Streamlit demo; then limitations and
        who did what.

        The headline number to remember: on 5,672 held-out tracks the Random Forest reaches {f3(METRICS['test']['macro_f1'])}
        Macro F1 and {pct(METRICS['test']['recall3'])} Recall@3, and we got there with song-level splitting so the score is not inflated.
    """)


def slide_2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, 2, "Problem formulation and practical scope", "PROBLEM")

    gap = 0.3
    w = (CONTENT_W - 2 * gap) / 3
    y, h = 1.5, 5.3
    cards = [
        ("Baseline and gap", [
            {"lead": "Baseline: ", "text": "Stanford CS229 work treats playlist assignment as supervised learning on Spotify audio features.", "space_after": 10},
            {"lead": "Gap: ", "text": "tracks belong to several genres, and the same track appears in several playlists.", "space_after": 10},
            {"lead": "Our focus: ", "text": "multi-label targets and an evaluation that cannot see the same song in train and test."},
        ]),
        ("Content-based proxy", [
            {"text": "Input: 12 acoustic attributes of one track.", "bullet": True, "space_after": 8},
            {"text": "Output: a probability for each of 6 genres, ranked.", "bullet": True, "space_after": 8},
            {"text": "The curator suggests the top-ranked genres as playlist homes.", "bullet": True, "space_after": 8},
            {"text": "No listening history and no user profile are used.", "bullet": True},
        ]),
        ("Honest scope", [
            {"text": "A classifier trained on an archival CSV, not a full recommender.", "bullet": True, "space_after": 8},
            {"text": "Playlist metadata and track popularity are excluded on purpose.", "bullet": True, "space_after": 8},
            {"text": "Results describe this dataset and these six genres only.", "bullet": True, "space_after": 8},
            {"text": "The demo ranks tracks from a CSV and makes no Spotify API calls.", "bullet": True},
        ]),
    ]
    for i, (title, body) in enumerate(cards):
        x = MARGIN + i * (w + gap)
        _card(s, x, y, w, h)
        _number_badge(s, x + 0.25, y + 0.3, i + 1)
        _text(s, x + 0.85, y + 0.3, w - 1.1, 0.46, title, size=18, color=NAVY, bold=True,
              anchor=MSO_ANCHOR.MIDDLE, label=f"s2 title {i}")
        _rect(s, x + 0.25, y + 1.0, w - 0.5, 0.02, fill=BORDER)
        _text(s, x + 0.25, y + 1.25, w - 0.5, h - 1.5, body, size=16, color=CHARCOAL, label=f"s2 body {i}")

    _notes(s, """
        Problem framing. The task: from the audio features of a track, predict which playlist genres it fits.

        Card 1: We frame the work against the Stanford CS229 line of playlist-genre classification on Spotify features.
        Our additions are multi-label targets (one track can sit in more than one genre) and strict song-level splitting.
        (Check the exact wording of the CS229 reference against our written report before presenting.)

        Card 2: Our curator is content-based. It sees only acoustic attributes, outputs six probabilities, and ranks genres.
        It has no notion of a listener.

        Card 3: Be upfront about scope. This is a classifier on an archival dataset. It is not a production recommender,
        and the demo does not call Spotify.
    """)


def slide_3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, 3, "Dataset and multi-label ground truth", "DATA")

    stats = [("32,833", "playlist-track rows"), ("28,356", "unique tracks"),
             ("12", "acoustic features"), ("6", "genre labels")]
    gap = 0.25
    w = (CONTENT_W - 3 * gap) / 4
    for i, (val, label) in enumerate(stats):
        x = MARGIN + i * (w + gap)
        _card(s, x, 1.5, w, 1.35, fill=MINT, line=TEAL)
        _text(s, x, 1.65, w, 0.65, val, size=34, color=TEAL, bold=True, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE, label=f"stat {i}")
        _text(s, x, 2.35, w, 0.3, label, size=13, color=SLATE, align=PP_ALIGN.CENTER, label=f"stat label {i}")

    y, h = 3.1, 3.7
    w2 = (CONTENT_W - 0.3) / 2
    # left card
    _card(s, MARGIN, y, w2, h)
    _text(s, MARGIN + 0.3, y + 0.25, w2 - 0.6, 0.4, "12 acoustic features", size=18, color=NAVY, bold=True, label="s3 left title")
    _text(s, MARGIN + 0.3, y + 0.8, w2 - 0.6, 1.1,
          "danceability, energy, key, loudness, mode, speechiness, acousticness, instrumentalness, liveness, valence, tempo, duration_ms",
          size=14, color=CHARCOAL, label="s3 features")
    _rect(s, MARGIN + 0.3, y + 2.1, w2 - 0.6, 0.02, fill=BORDER)
    _text(s, MARGIN + 0.3, y + 2.3, w2 - 0.6, 1.2, [
        {"lead": "Zero-leakage guarantee: ", "lead_color": TEAL,
         "text": "playlist metadata and track popularity never enter the feature matrix, so the model cannot read the answer from the playlist it came from."},
    ], size=14, label="s3 leakage")

    # right card
    x2 = MARGIN + w2 + 0.3
    _card(s, x2, y, w2, h)
    _text(s, x2 + 0.3, y + 0.25, w2 - 0.6, 0.4, "Multi-label ground truth", size=18, color=NAVY, bold=True, label="s3 right title")
    _text(s, x2 + 0.3, y + 0.8, w2 - 0.6, 2.8, [
        {"text": "Genres: edm, latin, pop, r&b, rap, rock.", "bullet": True, "space_after": 8},
        {"text": "Each track gets 6 binary flags, so one song can sit in several playlists.", "bullet": True, "space_after": 8},
        {"lead": "1,687 crossover tracks ", "lead_color": TEAL,
         "text": "(about 5.9% of 28,356) carry more than one genre label.", "bullet": True, "space_after": 8},
        {"text": "32,833 rows collapse to 28,356 tracks by track_id.", "bullet": True},
    ], size=14, label="s3 right body")

    _notes(s, """
        Dataset. An archival Spotify songs CSV with 32,833 playlist-track rows. After deduplicating by track_id we have
        28,356 unique tracks, because the same song can appear under more than one playlist.

        Features: 12 acoustic attributes (danceability through duration_ms). We deliberately left out playlist metadata
        and track popularity, since both would hand the model information about the label.

        Target: six genres as binary flags. 1,687 tracks (about 5.9%) have more than one genre. That is why we use
        multi-label metrics and report exact subset accuracy, which requires every one of the six flags to be right.
    """)


def slide_4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, 4, "Leakage prevention and pipeline design", "METHOD")

    steps = [
        ("Track-level aggregation", "Rows are merged by track_id, so each song is one sample with its full genre set."),
        ("Strict 80/20 song-level split", "Grouped by track_id before splitting, seed 42: 22,684 train and 5,672 test tracks."),
        ("Encapsulated 5-fold CV", "Median imputation and standard scaling sit inside a sklearn Pipeline, refit per fold."),
        ("Macro F1 selection", "The model with the best mean CV Macro F1 is chosen, so small genres count equally."),
    ]
    gap = 0.25
    w = (CONTENT_W - 3 * gap) / 4
    y, h = 1.5, 2.85
    for i, (title, body) in enumerate(steps):
        x = MARGIN + i * (w + gap)
        _card(s, x, y, w, h)
        _number_badge(s, x + 0.22, y + 0.22, i + 1)
        _text(s, x + 0.22, y + 0.85, w - 0.44, 0.65, title, size=15, color=NAVY, bold=True, label=f"s4 step title {i}")
        _text(s, x + 0.22, y + 1.55, w - 0.44, 1.2, body, size=14, color=CHARCOAL, label=f"s4 step body {i}")

    by = 4.6
    bh = 2.2
    _card(s, MARGIN, by, CONTENT_W, bh, fill=MINT, line=TEAL)
    _text(s, MARGIN + 0.35, by + 0.25, CONTENT_W - 0.7, 0.4, "Why row-level splitting gives optimistic scores",
          size=18, color=TEAL, bold=True, label="s4 callout title")
    _text(s, MARGIN + 0.35, by + 0.8, CONTENT_W - 0.7, 1.3, [
        {"text": "The 4,477 extra rows (32,833 minus 28,356) are repeat appearances of the same track.", "bullet": True, "space_after": 6},
        {"text": "A random row split puts copies of one song in both train and test, so the model is graded on songs it has already seen.", "bullet": True, "space_after": 6},
        {"text": "Splitting by track_id first keeps every song on exactly one side, so test scores reflect unseen music.", "bullet": True},
    ], size=14, label="s4 callout body")

    _notes(s, """
        Leakage prevention is the most important methodological choice.

        Step 1: aggregate to one row per track. Step 2: split 80/20 by track_id before anything else, seed 42, giving
        22,684 training and 5,672 test tracks. Step 3: run 5-fold cross-validation on the development set only, with
        median imputation and standard scaling inside a Pipeline so each fold fits those statistics on its own training
        portion. Step 4: pick the model by mean CV Macro F1.

        Why it matters: 32,833 rows versus 28,356 tracks means 4,477 repeat rows. With a row-level split those repeats
        land on both sides and inflate the score. We avoid that.
    """)


def slide_5(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, 5, "5-fold cross-validation model comparison", "MODEL SELECTION")

    cv = METRICS["cv"]
    header = ["Model", "Micro F1", "Macro F1", "Top-1 hit", "Recall@3", "Subset acc."]
    rows = [header]
    for m in MODELS:
        d = cv[m]
        name = "Random Forest (100 trees)" if m == "Random Forest" else m
        rows.append([name, f3(d["micro_f1"]), f3(d["macro_f1"]), pct(d["top1"]), pct(d["recall3"]), pct(d["subset"])])

    tw = 8.3
    col_w = [2.55, 1.15, 1.15, 1.15, 1.15, 1.15]
    assert abs(sum(col_w) - tw) < 1e-6
    _table(s, MARGIN, 1.5, col_w, [0.6, 0.8, 0.8, 0.8], rows, size=14, header_size=13, highlight_rows=(3,))

    # takeaway card under the table (table ends at 1.5 + 3.0 = 4.5)
    ty, th = 4.75, 2.05
    _card(s, MARGIN, ty, tw, th)
    d_svm = cv["Random Forest"]["macro_f1"] - cv["RBF SVM"]["macro_f1"]
    d_lr = cv["Random Forest"]["macro_f1"] - cv["Logistic Regression"]["macro_f1"]
    _text(s, MARGIN + 0.3, ty + 0.2, tw - 0.6, 0.35, "Reading the table", size=16, color=NAVY, bold=True, label="s5 read title")
    _text(s, MARGIN + 0.3, ty + 0.65, tw - 0.6, 1.3, [
        {"text": f"Mean over 5 folds on 22,684 development tracks. Random Forest leads Macro F1 by {d_svm:.3f} over RBF SVM and {d_lr:.3f} over Logistic Regression.", "bullet": True, "space_after": 6},
        {"text": "Top-1 and Recall@3 judge the ranking. Subset accuracy needs all 6 flags right, so it is the strictest column.", "bullet": True},
    ], size=13, label="s5 read body")

    cx = MARGIN + tw + 0.3
    cw = CONTENT_W - tw - 0.3
    _card(s, cx, 1.5, cw, 5.3, fill=MINT, line=TEAL)
    _text(s, cx + 0.3, 1.75, cw - 0.6, 0.7, "Why Random Forest won", size=18, color=TEAL, bold=True, label="s5 call title")
    _text(s, cx + 0.3, 2.5, cw - 0.6, 4.1, [
        {"text": f"Best mean Macro F1 ({f3(cv['Random Forest']['macro_f1'])}), the selection rule, and best on every column.", "bullet": True, "space_after": 8},
        {"text": "Trees can model interactions between features such as energy, loudness and acousticness.", "bullet": True, "space_after": 8},
        {"lead": "Logistic Regression: ", "text": f"subset accuracy is {pct(cv['Logistic Regression']['subset'])} against {pct(cv['Random Forest']['subset'])}. A linear boundary per genre likely struggles to get all six flags right on overlapping genres.", "bullet": True},
    ], size=14, label="s5 call body")

    _notes(s, f"""
        Model comparison, 5-fold CV on the development set.

        Logistic Regression: Micro F1 {f3(cv['Logistic Regression']['micro_f1'])}, Macro F1 {f3(cv['Logistic Regression']['macro_f1'])}, subset accuracy only {pct(cv['Logistic Regression']['subset'])}.
        RBF SVM: Micro F1 {f3(cv['RBF SVM']['micro_f1'])}, Macro F1 {f3(cv['RBF SVM']['macro_f1'])}, subset {pct(cv['RBF SVM']['subset'])}.
        Random Forest (100 trees): Micro F1 {f3(cv['Random Forest']['micro_f1'])}, Macro F1 {f3(cv['Random Forest']['macro_f1'])}, subset {pct(cv['Random Forest']['subset'])}.

        Random Forest was selected on mean CV Macro F1. The explanation for the linear model's low subset accuracy is our
        interpretation (one linear boundary per genre, overlapping classes), not something we tested directly. Say so if asked.
    """)


def slide_6(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, 6, "Final held-out test evaluation", "RESULTS")

    t = METRICS["test"]
    pills = [
        (f3(t["macro_f1"]), "Macro F1"),
        (f3(t["micro_f1"]), "Micro F1"),
        (pct(t["subset"]), "Exact subset acc."),
        (pct(t["top1"]), "Top-1 genre hit"),
        (pct(t["recall3"]), "Recall@3"),
    ]
    gap = 0.2
    w = (CONTENT_W - 4 * gap) / 5
    for i, (val, label) in enumerate(pills):
        x = MARGIN + i * (w + gap)
        _card(s, x, 1.5, w, 1.35, fill=TEAL, line=TEAL)
        _text(s, x, 1.62, w, 0.65, val, size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE, label=f"pill {i}")
        _text(s, x, 2.35, w, 0.3, label, size=12, color=MINT, align=PP_ALIGN.CENTER, label=f"pill label {i}")

    y, h = 3.1, 3.7
    lw = 6.6
    _card(s, MARGIN, y, lw, h)
    _text(s, MARGIN + 0.3, y + 0.2, lw - 0.6, 0.35, "Per-genre test F1 (5,672 held-out tracks)", size=16, color=NAVY, bold=True, label="s6 table title")
    genre_rows = [["Genre", "Test F1", "Test tracks"]]
    for name, f1, n in METRICS["per_genre"]:
        genre_rows.append([name, f3(f1), f"{n:,}"])
    _table(s, MARGIN + 0.3, y + 0.7, [2.4, 1.8, 1.8], [0.42] * 7, genre_rows, size=13, header_size=12)

    rx = MARGIN + lw + 0.3
    rw = CONTENT_W - lw - 0.3
    _card(s, rx, y, rw, h, fill=MINT, line=TEAL)
    _text(s, rx + 0.3, y + 0.2, rw - 0.6, 0.4, "Why Rock and EDM beat Pop", size=18, color=TEAL, bold=True, label="s6 call title")
    _text(s, rx + 0.3, y + 0.8, rw - 0.6, 2.8, [
        {"text": "Rock (0.703) and EDM (0.683) have distinctive acoustic profiles in energy, loudness and tempo.", "bullet": True, "space_after": 8},
        {"text": "Pop (0.391) is an umbrella genre that overlaps with all the others, so no feature region is clearly its own.", "bullet": True, "space_after": 8},
        {"text": "This is our likely explanation. We did not run an ablation to confirm it.", "bullet": True},
    ], size=15, label="s6 call body")

    _notes(s, f"""
        Final test evaluation of the Random Forest on 5,672 tracks never used for training or model selection.

        Macro F1 {f3(t['macro_f1'])}, Micro F1 {f3(t['micro_f1'])}, exact subset accuracy {pct(t['subset'], 2)}, top-1 genre hit {pct(t['top1'], 2)}, Recall@3 {pct(t['recall3'], 2)}.
        These sit close to the CV numbers (Macro F1 0.556), which is what we expect when leakage is controlled.

        Per genre: Rock 0.703, EDM 0.683, Rap 0.653, R&B 0.491, Latin 0.457, Pop 0.391.
        The reason Pop scores lowest is our hypothesis: it is a broad label that overlaps the rest. We have not tested it.
    """)


def slide_7(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, 7, "System architecture and live demo", "DEMO")

    blocks = [
        ("Model bundle", "results/playlist_model.joblib"),
        ("CSV ingestion", "samples/demo_tracks.csv"),
        ("Probabilistic ranking", "Genre probabilities per track, sorted high to low"),
        ("Streamlit UI + CSV", "app.py with a downloadable ranked CSV"),
    ]
    arrow_gap = 0.5
    w = (CONTENT_W - 3 * arrow_gap) / 4
    y, h = 1.5, 2.2
    for i, (title, sub) in enumerate(blocks):
        x = MARGIN + i * (w + arrow_gap)
        _card(s, x, y, w, h)
        _number_badge(s, x + 0.22, y + 0.22, i + 1, d=0.42)
        _text(s, x + 0.22, y + 0.8, w - 0.44, 0.35, title, size=14, color=NAVY, bold=True, label=f"s7 block title {i}")
        is_path = i < 2
        _text(s, x + 0.22, y + 1.25, w - 0.44, 0.8, sub, size=12 if is_path else 14,
              color=TEAL if is_path else CHARCOAL, font=MONO if is_path else FONT, label=f"s7 block sub {i}")
        if i < 3:
            _rect(s, x + w + 0.1, y + h / 2 - 0.15, arrow_gap - 0.2, 0.3, fill=TEAL, shape=MSO_SHAPE.RIGHT_ARROW)

    cy, ch = 4.0, 2.8
    _card(s, MARGIN, cy, CONTENT_W, ch)
    _text(s, MARGIN + 0.35, cy + 0.25, CONTENT_W - 0.7, 0.4, "Live demo verification checklist", size=18, color=NAVY, bold=True, label="s7 check title")
    left = ["□  Model bundle loads without retraining",
            "□  demo_tracks.csv uploads with the 12 feature columns",
            "□  Ranked genres appear with probabilities"]
    right = ["□  Top-3 genres shown for each track",
             "□  Download button returns a ranked CSV",
             "□  Runs with no Spotify API key or network call"]
    colw = (CONTENT_W - 0.7 - 0.4) / 2
    _text(s, MARGIN + 0.35, cy + 0.9, colw, 1.7, left, size=16, space_after=14, label="s7 check left")
    _text(s, MARGIN + 0.35 + colw + 0.4, cy + 0.9, colw, 1.7, right, size=16, space_after=14, label="s7 check right")

    _notes(s, """
        Demo flow. The trained Random Forest is saved as results/playlist_model.joblib. The Streamlit app (app.py) loads it,
        reads a CSV such as samples/demo_tracks.csv, computes a probability for each of the six genres per track, ranks them,
        and shows the result in the UI with a download button for the ranked CSV.

        Before presenting, run through the checklist on the slide: model loads, CSV uploads, ranking shows, top-3 appears,
        download works, and nothing needs internet or a Spotify key. Keep the CSV open as a fallback if the live app fails.
    """)


def slide_8(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, 8, "Limitations and individual contributions", "WRAP-UP")

    y, h = 1.5, 5.3
    lw = 5.4
    _card(s, MARGIN, y, lw, h)
    _text(s, MARGIN + 0.3, y + 0.25, lw - 0.6, 0.4, "Project limitations", size=18, color=NAVY, bold=True, label="s8 left title")
    _rect(s, MARGIN + 0.3, y + 0.85, lw - 0.6, 0.02, fill=BORDER)
    _text(s, MARGIN + 0.3, y + 1.1, lw - 0.6, 4.0, [
        {"lead": "Proxy, not taste: ", "text": "genre labels stand in for a curator. No listener preference data is used.", "bullet": True, "space_after": 10},
        {"lead": "Historical dataset: ", "text": "an archival CSV, so features and tags may not match today's catalog.", "bullet": True, "space_after": 10},
        {"lead": "No Spotify API: ", "text": "the demo scores only tracks supplied in a CSV.", "bullet": True, "space_after": 10},
        {"lead": "Noisy labels: ", "text": "playlist-derived genres overlap, and Pop reaches only 0.391 F1.", "bullet": True, "space_after": 10},
        {"lead": "One split: ", "text": "a single 80/20 holdout with seed 42.", "bullet": True},
    ], size=15, label="s8 left body")

    rx = MARGIN + lw + 0.3
    rw = CONTENT_W - lw - 0.3
    _card(s, rx, y, rw, h)
    _text(s, rx + 0.3, y + 0.25, rw - 0.6, 0.4, "Individual contribution record", size=18, color=NAVY, bold=True, label="s8 right title")
    rows = [
        ["Member", "Contribution"],
        ["Mohammed Adaan Hamad",
         "Dataset preparation, CSV validation, track-level grouping, multi-label target formulation, model comparison, cross-validation scoring."],
        ["Lakshya Jeet Singh",
         "Evaluation methodology, train/test leakage prevention design, feature selection validation, Streamlit demo (app.py), automated report and presentation pipeline, final QA and testing."],
    ]
    _table(s, rx + 0.3, y + 1.0, [1.9, rw - 0.6 - 1.9], [0.5, 1.6, 1.9], rows, size=13, header_size=13, left_cols=(1,))

    _notes(s, """
        Limitations first. The curator is a proxy: genres are playlist-derived labels, not personal taste. The data is a
        historical snapshot. There is no Spotify API scraping, so the demo only scores tracks in a CSV. Labels overlap,
        which hurts Pop most. And we report a single 80/20 split.

        Contributions. Mohammed Adaan Hamad: dataset preparation, CSV validation, track-level grouping, multi-label target
        formulation, model comparison and cross-validation scoring. Lakshya Jeet Singh: evaluation methodology, leakage
        prevention design, feature selection validation, the Streamlit demo (app.py), the automated report and
        presentation pipeline, and final QA and testing.

        Close by restating the result: Random Forest, Macro F1 0.563 on held-out tracks, with leakage controlled.
    """)


# ---------------------------------------------------------------- Keynote compatibility
def _fix_keynote_compat(pptx_path):
    """python-pptx leaves <p:sldSz type="screen4x3"/> after a 16:9 resize; Keynote rejects that."""
    with open(pptx_path, 'rb') as f:
        in_zip = zipfile.ZipFile(f)
        out_buf = io.BytesIO()
        with zipfile.ZipFile(out_buf, 'w', compression=zipfile.ZIP_DEFLATED) as out:
            for item in in_zip.infolist():
                data = in_zip.read(item.filename)
                if item.filename == 'ppt/presentation.xml':
                    text = data.decode('utf-8')
                    text = re.sub(r'type="screen4x3"', 'type="screen16x9"', text)
                    data = text.encode('utf-8')
                out.writestr(item, data)
    with open(pptx_path, 'wb') as f:
        f.write(out_buf.getvalue())


# ---------------------------------------------------------------- main
def build():
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)

    for fn in (slide_1, slide_2, slide_3, slide_4, slide_5, slide_6, slide_7, slide_8):
        fn(prs)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PATH))
    _fix_keynote_compat(str(OUT_PATH))

    with zipfile.ZipFile(OUT_PATH) as z:
        xml = z.read("ppt/presentation.xml").decode("utf-8")
    m = re.search(r'<p:sldSz[^>]*>', xml)
    print(f"[ok] Saved {OUT_PATH}")
    print(f"[ok] {m.group(0) if m else 'sldSz not found'}")
    if WARNINGS:
        print("[warn] Layout warnings (estimates, check the rendered slides):")
        for w in WARNINGS:
            print("  -", w)
    else:
        print("[ok] No text-overflow warnings.")


if __name__ == "__main__":
    build()
